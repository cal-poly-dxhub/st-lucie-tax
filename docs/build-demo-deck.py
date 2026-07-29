#!/usr/bin/env python3
"""Generate the pre-demo slide deck for the St. Lucie chatbot.
Minimal styling on purpose: title + bullets only, default template."""

from pptx import Presentation
from pptx.util import Pt

prs = Presentation()

TITLE_SLIDE = prs.slide_layouts[0]   # title + subtitle
BULLET_SLIDE = prs.slide_layouts[1]  # title + content


def add_title(title, subtitle):
    s = prs.slides.add_slide(TITLE_SLIDE)
    s.shapes.title.text = title
    s.placeholders[1].text = subtitle
    return s


def add_bullets(title, bullets):
    s = prs.slides.add_slide(BULLET_SLIDE)
    s.shapes.title.text = title
    body = s.placeholders[1].text_frame
    body.clear()
    for i, b in enumerate(bullets):
        p = body.paragraphs[0] if i == 0 else body.add_paragraph()
        p.text = b
        p.font.size = Pt(18)
    return s


# --- Title ---
add_title(
    "St. Lucie County Chatbot",
    "Progress since the last demo  |  June 24, 2026",
)

# 1. Security fixes
add_bullets("1. Security Fixes", [
    "Completed external security review remediation (Caylent)",
    "Secrets moved out of code into encrypted AWS Secrets Manager",
    "Citizen PII (SSN, card, bank #) scrubbed before it reaches the AI",
    "Sensitive data redacted from logs and debug tools",
    "Logins now expire and auto-refresh; abuse blocked at two layers",
    "AI access locked to only the models and data we authorize",
])

# 2. Accuracy boost from SME testing
add_bullets("2. Accuracy Boost (SME Testing)", [
    "Driven by your tax-office staff testing real transactions",
    "Bot no longer invents phone numbers, fees, or policy",
    "Deceased-owner / inheritance title transfers now handled correctly",
    "First-time license applicants no longer hit circular blocks",
    "Redundant questions removed; out-of-state CDL routing fixed",
    "Official source links now surface in chat and side panel",
    "Every fix verified against Florida statute or staff feedback",
])

# 3. Mobile optimization
add_bullets("3. Mobile Optimization", [
    "Rebuilt mobile-first for QR-code, phone-majority traffic",
    "Everything fits on a phone: no overflow, no clipped text",
    "Slide-out menu, progress bar, animated quick-reply chips",
    "Identity verification never dead-ends; skip always works",
    "Stale verification tokens now refresh invisibly",
    "Staff tools: reviewed checkbox + inline transcript feedback",
    "Faster admin navigation with cached session list",
])

# 4. Production prep
add_bullets("4. Production Prep", [
    "Before-production security checklist documented",
    "Deployment runbook for secrets population in place",
    "Architecture teaching kit for technical partners",
    "Interactive explorer captured from the live system",
    "One-page cheat sheet and guided code tour for handoff",
    "Graceful fallbacks: scheduling and verify degrade safely",
    "Remaining items scoped: Cognito, VPC/KMS, deploy",
])

# 5. SME questions material prep
add_bullets("5. SME Question Materials", [
    "Policy-questions workbook: 46 questions across 3 tabs",
    "12 high-impact, prioritized for fastest review",
    "Yellow answer column, source citations, unverified flags",
    "Upload-review spreadsheet: 168 documents in 7 categories",
    "One-click checkboxes; Excel and Google Sheets copies",
    "Pulled from all decision trees + real tester sessions",
    "Each answer maps to a simple, one-line code change",
])

out = "demo-deck-2026-06-24.pptx"
prs.save(out)
print("wrote", out, "with", len(prs.slides._sldIdLst), "slides")
